import time

import requests
from selenium import webdriver
from selenium.webdriver import ActionChains
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.common.by import By
#
# inp = [12,24,13,54,67,23,25,37]
#
# prime_nos = []
# for i in inp:
#     for j in range(2,i):
#         if i%j == 0:
#             break
#     else:
#         prime_nos.append(i)
# print(prime_nos)


import openpyxl

# workbook = openpyxl.load_workbook("Data/practice_sheet.xlsx")
# sheet = workbook['Sheet1']
#
# print(sheet.cell(row=2,column=2).value)
#
# max_row = sheet.max_row
# max_column = sheet.max_column
# print(max_row)
# print(max_column)
#
# for column in sheet.iter_cols(min_row=1,max_col=max_column,max_row=max_row):
#     for cell in column:
#         cell_value = cell.value
#         print(cell_value)


# workbook.close()

service_obj = Service("Drivers/chromedriver_win32/chromedriver.exe")
options = webdriver.ChromeOptions()
options.add_argument("--headless")

driver = webdriver.Chrome(service=service_obj, options=options)
driver.get("https://automationexercise.com/")

time.sleep(2)

image = driver.find_element(By.XPATH, '/html/body/header/div/div/div/div[1]/div/a/img')

image_src_before_hover = image.get_attribute('src')

act = ActionChains(driver)
act.move_to_element(image).perform()

time.sleep(2)

image_src_after_hover = image.get_attribute('src')

assert image_src_before_hover == image_src_after_hover

from requests.exceptions import *

try:
    response = requests.post("https://httpbin.org/#/", params=None, data=None, verify=False)
    print(response.headers)
    print(response.status_code)

except RequestException as e:
    print(e)
